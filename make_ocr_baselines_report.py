#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""OCR 비교군 결과 표 생성 — GSV(eval_ocr_v2 summary) + in-domain(indomain_summary.csv).

입력:
  artifacts/gt/ocr_eval_v2_summary.csv  (eval_ocr_v2.py --engines ... 실행 후; region/engine 행)
  artifacts/str_baselines/indomain_summary.csv
출력:
  artifacts/str_baselines/ocr_baselines_tables.md   (README 붙여넣기용)
  Desktop/GSV_results_v7.docx                       (v5 복사 + Table 5/6 추가; v6 = Tesseract 미세조정 전)
Usage: .venv/Scripts/python.exe make_ocr_baselines_report.py [--engines easyocr,trocr,tesseract,surya,parseq,svtrv2,paddle] [--no-docx]
"""
import argparse
import csv
import shutil
from collections import defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
GSV = HERE / "artifacts" / "gt" / "ocr_eval_v2_summary.csv"
IND = HERE / "artifacts" / "str_baselines" / "indomain_summary.csv"
OUT_MD = HERE / "artifacts" / "str_baselines" / "ocr_baselines_tables.md"
DESK = Path.home() / "Desktop"
LABEL = {
    "tesseract": ("Tesseract 5.5 (kor+eng)", "general engine, off-the-shelf"),
    "tesseract_ft": ("Tesseract 5.5 (kor fine-tuned + eng)", "general engine, fine-tuned, same data"),
    "easyocr": ("EasyOCR (CRNN, fine-tuned v3)", "general engine, per-box"),
    "surya": ("Surya 0.14 (rec2)", "general engine, off-the-shelf"),
    "trocr": ("TrOCR-small (fine-tuned v3)", "document STR, per-box"),
    "parseq": ("PARSeq (ViT-S, fine-tuned)", "academic STR, same data"),
    "svtrv2": ("SVTRv2-B (fine-tuned)", "academic STR, same data"),
    "paddle1": ("PP-OCRv5 rec (v5_lines, single model)", "fine-tuned, same data, no vote"),
    "paddle": ("PP-OCRv5 rec (v5_lines, deployed: 3-way vote + en model)", "deployed pipeline"),
}
REGIONS = ["gangnam", "brooklyn", "suwon", "ALL"]
RNAME = {"gangnam": "Gangnam", "brooklyn": "Brooklyn", "suwon": "Suwon", "ALL": "Total"}


def read_gsv(engines):
    rows = defaultdict(dict)
    with GSV.open(encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if r["engine"] in engines:
                rows[r["engine"]][r["region"]] = r
    return rows


def read_ind(engines):
    rows = defaultdict(dict)
    if not IND.exists():
        return rows
    with IND.open(encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if r["engine"] in engines:
                rows[r["engine"]][r["set"]] = r     # 마지막 행이 이김(재실행 시)
    return rows


def cell(r):
    return f"{float(r['exact'])*100:.1f}% / {float(r['CER_micro']):.3f} / {float(r['WER']):.3f}" if r else "—"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--engines", default="tesseract,tesseract_ft,easyocr,surya,trocr,parseq,svtrv2,paddle1,paddle")
    ap.add_argument("--no-docx", action="store_true")
    args = ap.parse_args()
    engines = [e for e in args.engines.split(",") if e]
    gsv, ind = read_gsv(engines), read_ind(engines)

    md = ["### GSV 라인 채점 — line exact / CER / WER (eval_ocr_v2, --mask-phone, brooklyn en-only)", "",
          "| 엔진 | 부류 | " + " | ".join(RNAME[r] for r in REGIONS) + " | WAR (Total) |",
          "|---|---|" + "---|" * (len(REGIONS) + 1)]
    for e in engines:
        if e not in gsv:
            continue
        name, kind = LABEL.get(e, (e, ""))
        war = gsv[e].get("ALL", {}).get("WAR")
        md.append(f"| {name} | {kind} | " + " | ".join(cell(gsv[e].get(r)) for r in REGIONS)
                  + f" | {float(war):.3f} |" if war else f"| {name} | {kind} | " + " | ".join(cell(gsv[e].get(r)) for r in REGIONS) + " | — |")
    md += ["", "### In-domain (signboard_v3 test) — exact / CER / WER, 1:1 크롭 채점", "",
           "| 엔진 | test_line (1,631 실라인) | test_word (7,756 단어) |", "|---|---|---|"]
    for e in engines:
        if e not in ind:
            continue
        name, _ = LABEL.get(e, (e, ""))
        def c(r):
            return f"{float(r['exact'])*100:.1f}% / {float(r['CER']):.3f} / {float(r['WER']):.3f}" if r else "—"
        md.append(f"| {name} | {c(ind[e].get('test_line'))} | {c(ind[e].get('test_word'))} |")
    OUT_MD.write_text("\n".join(md) + "\n", encoding="utf-8")
    print("\n".join(md))

    if args.no_docx:
        return
    from docx import Document
    src, dst = DESK / "GSV_results_v5.docx", DESK / "GSV_results_v7.docx"
    shutil.copy(src, dst); d = Document(dst)
    d.add_paragraph("")
    d.add_paragraph("Table 5. OCR engines on GSV signboard crops by region (line exact / CER / WER)")
    d.add_paragraph("Same detection (CRAFT ∪ PaddleOCR-DB) and line merging for every engine; only the recognizer differs. "
                    "Scoring: eval_ocr_v2 line matching, NFKC + lowercase + [0-9a-z가-힣] filter, phone numbers masked, "
                    "Brooklyn scored English-only. WER = word-level edit distance / GT words (spurious lines count as insertions, so >1 is possible). "
                    "Fine-tuned STR models (PARSeq, SVTRv2, PP-OCRv5) use the identical training mix (62,464 word + 13,148 line crops, signboard_v3 split).")
    hdr = ["Engine", "Category"] + [RNAME[r] for r in REGIONS] + ["WAR (Total)"]
    t = d.add_table(rows=1, cols=len(hdr)); t.style = d.tables[0].style
    for j, v in enumerate(hdr):
        t.cell(0, j).text = v
        for run in t.cell(0, j).paragraphs[0].runs: run.bold = True
    for e in engines:
        if e not in gsv:
            continue
        name, kind = LABEL.get(e, (e, ""))
        war = gsv[e].get("ALL", {}).get("WAR")
        vals = [name, kind] + [cell(gsv[e].get(r)) for r in REGIONS] + [f"{float(war):.3f}" if war else "—"]
        row = t.add_row().cells
        for j, v in enumerate(vals): row[j].text = v
    if ind:
        d.add_paragraph("")
        d.add_paragraph("Table 6. OCR recognizers on the in-domain signboard_v3 test split (exact / CER / WER)")
        d.add_paragraph("1:1 crop scoring (no line matching). test_line = 1,631 real multi-word line crops; test_word = 7,756 word crops.")
        hdr = ["Engine", "test_line (1,631)", "test_word (7,756)"]
        t = d.add_table(rows=1, cols=3); t.style = d.tables[0].style
        for j, v in enumerate(hdr):
            t.cell(0, j).text = v
            for run in t.cell(0, j).paragraphs[0].runs: run.bold = True
        for e in engines:
            if e not in ind:
                continue
            name, _ = LABEL.get(e, (e, ""))
            def c(r):
                return f"{float(r['exact'])*100:.1f}% / {float(r['CER']):.3f} / {float(r['WER']):.3f}" if r else "—"
            row = t.add_row().cells
            for j, v in enumerate([name, c(ind[e].get("test_line")), c(ind[e].get("test_word"))]): row[j].text = v
    d.save(dst); print("saved", dst)


if __name__ == "__main__":
    main()
